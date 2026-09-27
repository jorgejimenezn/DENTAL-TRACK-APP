"""Convierte el Excel de datos simulados en el archivo de carga del backend.

Uso (desde la raíz del repositorio):
    pip install openpyxl bcrypt
    python herramientas/preparar_datos_prueba.py datos_simulados_de_prueba.xlsx

Genera:
  - backend/app/db/datos_prueba.json   → se sube al repositorio (solo contiene hashes, no contraseñas)
  - credenciales_prueba.xlsx           → NO se sube; contiene las contraseñas iniciales

Transformaciones aplicadas:
  - NIT: se conserva el número y se recalcula el dígito verificador (los del Excel son ficticios).
  - Teléfonos de clientes (2xxx-xxxx) se guardan como teléfono fijo.
  - "Prioridad 2" = VIP, "Prioridad 3" = normal. La prioridad 1 queda para emergencias.
  - Departamento y municipio se deducen del texto de la dirección.
  - Rol y áreas de los empleados según la distribución propuesta.
"""
import json
import re
import secrets
import string
import sys
from pathlib import Path

import bcrypt
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

RAIZ = Path(__file__).resolve().parent.parent
SALIDA_JSON = RAIZ / "backend" / "app" / "db" / "datos_prueba.json"
SALIDA_CREDENCIALES = RAIZ / "credenciales_prueba.xlsx"

CORRECCIONES_CORREO = {"a.guilar@innovalab.com.gt": "a.aguilar@innovalab.com.gt"}

# rol, áreas
DISTRIBUCION = {
    "Jorge Jimenez": ("admin", []),
    "Axel Aguilar": ("admin", []),
    "María José Hernández": ("ingresos", ["Ingresos"]),
    "Ana Sofía Castillo": ("ingresos", ["Ingresos"]),
    "Roberto Estrada": ("logistica", ["Logística"]),
    "Mario Alberto Ruano": ("logistica", ["Logística"]),
    "Francisco Cifuentes": ("logistica", ["Logística"]),
    "Claudia Patricia Ramos": ("qa", ["QA"]),
    "Gabriela Michelle Soto": ("qa", ["QA"]),
    "Sofía Alejandra Flores": ("facturacion", ["Facturación_Cobros"]),
    "Daniela Andrea Gutiérrez": ("facturacion", ["Facturación_Cobros"]),
    "Carlos Mendoza": ("tecnico", ["Yesos"]),
    "Juan Carlos Morales": ("tecnico", ["Yesos"]),
    "Luis Fernando Gómez": ("tecnico", ["Diseño"]),
    "Diego Alejandro Ortiz": ("tecnico", ["Diseño", "CAD"]),
    "José Manuel Castro": ("tecnico", ["CAD"]),
    "Eduardo Ramírez": ("tecnico", ["CAD"]),
    "Andrea Lucía Monzón": ("tecnico", ["CAD"]),
    "Ricardo Enrique López": ("tecnico", ["CAM_IMP"]),
    "Sergio Armando Pineda": ("tecnico", ["CAM_IMP"]),
    "Alejandro Morales": ("tecnico", ["CAM_FRES"]),
    "Gabriel Antonio Juárez": ("tecnico", ["CAM_FRES", "CAM_IMP"]),
    "Óscar David Hernández": ("tecnico", ["Finalización"]),
    "Héctor Julio Mejía": ("tecnico", ["Finalización"]),
    "Pablo Andrés Vásquez": ("tecnico", ["Finalización", "Yesos"]),
}


def limpiar(valor) -> str:
    return re.sub(r"[\s\-]", "", str(valor or "")).upper()


def nit_con_verificador(nit: str) -> str:
    cuerpo = limpiar(nit)[:-1]
    total = sum(int(d) * (i + 2) for i, d in enumerate(reversed(cuerpo)))
    calculado = (11 - total % 11) % 11
    return f"{cuerpo}-{'K' if calculado == 10 else calculado}"


def separar_nombre(completo: str) -> tuple[str, str]:
    partes = completo.split()
    if len(partes) <= 2:
        return partes[0], " ".join(partes[1:])
    return " ".join(partes[:-1]), partes[-1]


def ubicar(direccion: str) -> tuple[str, str]:
    texto = direccion.lower()
    if "mixco" in texto:
        return "Guatemala", "Mixco"
    if "carretera a el salvador" in texto or "santa catarina pinula" in texto:
        return "Guatemala", "Santa Catarina Pinula"
    return "Guatemala", "Guatemala"


def nueva_password() -> str:
    alfabeto = "".join(c for c in string.ascii_letters + string.digits if c not in "O0Il1")
    return "".join(secrets.choice(alfabeto) for _ in range(10))


def filas(hoja):
    for fila in hoja.iter_rows(min_row=3, values_only=True):
        valores = fila[1:]
        if any(v not in (None, "") for v in valores):
            yield valores


def main(ruta_excel: str):
    libro = openpyxl.load_workbook(ruta_excel, data_only=True)
    empleados, clientes, credenciales_emp, credenciales_cli = [], [], [], []

    for nombre_completo, correo, direccion, telefono in filas(libro["empleados"]):
        nombre_completo = str(nombre_completo).strip()
        if nombre_completo not in DISTRIBUCION:
            raise SystemExit(f"Falta rol para: {nombre_completo}")
        correo = CORRECCIONES_CORREO.get(correo.strip().lower(), correo.strip().lower())
        nombre, apellido = separar_nombre(nombre_completo)
        rol, areas = DISTRIBUCION[nombre_completo]
        password = nueva_password()
        empleados.append({
            "nombre": nombre, "apellido": apellido, "email": correo,
            "telefono": limpiar(telefono) or None, "direccion": (direccion or None),
            "rol": rol, "areas": areas,
            "password_hash": bcrypt.hashpw(password.encode(), bcrypt.gensalt()).decode(),
        })
        credenciales_emp.append([nombre_completo, correo, rol, ", ".join(areas) or "—", password])

    for razon, facturacion, nit, telefono, correo, direccion, prioridad in filas(libro["clientes"]):
        nivel = int(re.search(r"\d", str(prioridad)).group())
        nit_corregido = nit_con_verificador(nit)
        departamento, municipio = ubicar(direccion)
        password = nueva_password()
        clientes.append({
            "razon_social": razon.strip(), "nombre_facturacion": facturacion.strip(),
            "email": correo.strip().lower(), "telefono_fijo": limpiar(telefono),
            "tipo_documento": "NIT", "numero_documento": limpiar(nit_corregido),
            "nivel_prioridad": nivel,
            "direccion": {"departamento": departamento, "municipio": municipio, "direccion": direccion.strip()},
            "password_hash": bcrypt.hashpw(password.encode(), bcrypt.gensalt()).decode(),
        })
        credenciales_cli.append([razon.strip(), correo.strip().lower(), nit, nit_corregido,
                                 "2 · VIP" if nivel == 2 else "3 · Normal", password])

    SALIDA_JSON.write_text(json.dumps({"empleados": empleados, "clientes": clientes},
                                      ensure_ascii=False, indent=1), encoding="utf-8")

    salida = openpyxl.Workbook()
    hojas = [
        ("Empleados", ["Nombre", "Correo (usuario)", "Rol", "Áreas", "Contraseña inicial"], credenciales_emp),
        ("Clientes", ["Cliente / razón social", "Correo (usuario)", "NIT original", "NIT cargado",
                      "Prioridad", "Contraseña inicial"], credenciales_cli),
    ]
    for i, (titulo, encabezados, datos) in enumerate(hojas):
        hoja = salida.active if i == 0 else salida.create_sheet()
        hoja.title = titulo
        hoja.append(encabezados)
        for celda in hoja[1]:
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor="111111")
            celda.alignment = Alignment(vertical="center")
        for fila in datos:
            hoja.append(fila)
        for columna in hoja.columns:
            ancho = max(len(str(c.value or "")) for c in columna) + 3
            hoja.column_dimensions[columna[0].column_letter].width = min(ancho, 48)
        hoja.freeze_panes = "A2"
    salida.save(SALIDA_CREDENCIALES)
    print(f"{len(empleados)} empleados y {len(clientes)} clientes → {SALIDA_JSON.relative_to(RAIZ)}")
    print(f"Credenciales → {SALIDA_CREDENCIALES.name} (no lo subas al repositorio)")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "datos_simulados_de_prueba.xlsx")
